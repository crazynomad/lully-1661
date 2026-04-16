"""Build the round-2 follow-up .docx for client review.

Short, focused on the 7 follow-ups that came out of round 1. Thank-you +
what-we-learned up top, then a fillable amber table.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "requirements-round2-CLIENT-REVIEW.docx"


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    return h


def add_para(doc, text, bold=False, italic=False, size=11, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build():
    doc = Document()

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("lully 1661 — Website Requirements (Round 2)")
    r.bold = True
    r.font.size = Pt(20)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("A short follow-up based on your Round 1 answers")
    rs.italic = True
    rs.font.size = Pt(12)
    rs.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run("Prepared 2026-04-16")
    rm.font.size = Pt(10)
    rm.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    # Thanks + what we learned
    add_heading(doc, "Thank you", level=1)
    add_para(
        doc,
        "Your Round 1 answers unblocked most of the work. Ten of fifteen questions "
        "are fully resolved, and we have started on the scope & sitemap v1 and the "
        "platform recommendation in parallel.",
    )

    add_heading(doc, "Already resolved from Round 1", level=2)
    resolved = [
        "Brand spelling: lully 1661 (double L). Repository will be renamed.",
        "Languages: PT + EN, PT default, /pt/ and /en/ subpaths.",
        "Payments: Multibanco + credit cards (Visa / Mastercard).",
        "Reservations: The Fork, with email fallback.",
        "B2B: a showcase page with an email enquiry form.",
        "Recruitment: simple email-based applications.",
        "Brand guidelines: none formal yet — we will propose during design.",
        "Vector logos: received (black + grey .ai files).",
        "Copy and imagery: supplied by you.",
        "Ambient audio: user-triggered, off during shop/checkout.",
    ]
    for item in resolved:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item).font.size = Pt(11)

    add_heading(doc, "What we took on ourselves", level=2)
    add_para(
        doc,
        "On Q3 (platform) you left the choice to us. We will come back with a "
        "recommendation (Shopify vs. headless) alongside the sitemap v1 — you "
        "do not need to answer that here.",
    )

    doc.add_page_break()

    # Round 2 questions table
    add_heading(doc, "Round 2 — we need a bit more detail", level=1)
    add_para(
        doc,
        "Seven short questions, most of them came up from your Round 1 answers. "
        "Amber cells below are waiting for your answer.",
        italic=True,
    )
    doc.add_paragraph()

    rows = [
        ("#", "Question", "Your answer / decision"),
        (
            "F1",
            "Stores — you mentioned three Lisbon outlets open Tue–Sun. "
            "Please share each store’s address, neighborhood, and opening hours. "
            "Is there a “flagship” we should feature first?",
            "",
        ),
        (
            "F2",
            "Go-live target — “ASAP” helps us prioritize, but to plan we need a target month. "
            "Is there a press moment, anniversary, or campaign the site should launch alongside? "
            "(e.g. June 2026, July 2026, before summer holidays, etc.)",
            "",
        ),
        (
            "F3",
            "Click & collect cutoff — what is the latest order time for same-day pickup? "
            "Typical bakery pattern is “order before 14:00 for same-day pickup, "
            "after that it becomes next-day.” Do breads need an earlier cutoff than pastries?",
            "",
        ),
        (
            "F4",
            "Your Q15 answer describes a section for festive products + events you organize or supply. "
            "Your original brief also has a “Special Orders” section. "
            "We read these as the same thing and propose merging into "
            "“Festive & Bespoke Orders”. Please confirm or correct.",
            "",
        ),
        (
            "F5",
            "Vector versions of the three Baroque sub-marks (Meunier / Pâtissière / Caffetier) — "
            "are there .ai / .eps / .svg files? We only have JPEG raster files today, "
            "which limits how flexibly we can compose them in the Gareth × Sezin flyer style.",
            "",
        ),
        (
            "F6",
            "Existing web presence — any current website, Instagram, or press coverage "
            "we should migrate content from, link from the new site, or cite as “as seen in”?",
            "",
        ),
        (
            "F7",
            "B2B partners — any existing chefs / brands you’d like us to name on the "
            "“Lully Inside” page at launch? Or should that page launch as descriptive-only "
            "until first deals close?",
            "",
        ),
    ]

    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Light Grid Accent 1"
    widths = [Cm(1.6), Cm(10.5), Cm(5.5)]
    for row_idx, row_data in enumerate(rows):
        for col_idx, value in enumerate(row_data):
            cell = table.rows[row_idx].cells[col_idx]
            cell.width = widths[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(value)
            run.font.size = Pt(10)
            if row_idx == 0:
                run.bold = True
                shade_cell(cell, "1A1A1A")
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                shade_cell(cell, "FFF8E1")

    doc.add_paragraph()
    add_para(
        doc,
        "Reply in the cells above — this Google Doc is editable. Once you are done we will "
        "move to wireframes and a concrete phase-1 scope proposal.",
        italic=True,
    )
    doc.add_paragraph()
    add_para(doc, "Thank you.", italic=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
